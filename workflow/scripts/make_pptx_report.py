#!/usr/bin/env python3
"""
Build a minimal summary PPTX report:
  1. Clinical descriptive statistics (continuous + categorical, side by side)
  2. Clinical descriptive statistics stratified by the target variable
  3. Metabolomics summary heatmap
  4. Proteomics summary heatmap
  5. Transcriptomics summary heatmap
"""

import pandas as pd
from pathlib import Path
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

CONTINUOUS  = "results/clinics/descriptive/continuous.csv"
CATEGORICAL = "results/clinics/descriptive/categorical.csv"
STRATIFIED  = "results/clinics/descriptive_stratified/descriptive_table_stratified.csv"
HEATMAP_METABOLOMICS   = "results/metabolomics/summary_heatmap.png"
HEATMAP_PROTEOMICS     = "results/proteomics/summary_heatmap.png"
HEATMAP_TRANSCRIPTOMICS = "results/transcriptomics/summary_heatmap.png"
OUT = Path("results/report/atheromics_tvas_report.pptx")

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)


def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.6))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.bold = True
    run.font.size = Pt(22)


def set_cell(cell, text, bold=False, size=Pt(7)):
    cell.text = ""
    cell.margin_top = Pt(1)
    cell.margin_bottom = Pt(1)
    cell.margin_left = Pt(3)
    cell.margin_right = Pt(3)
    cell.vertical_anchor = 3  # MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "" if pd.isna(text) else str(text)
    run.font.bold = bold
    run.font.size = size


def add_df_table(slide, df, left, top, width, height):
    n_rows, n_cols = df.shape[0] + 1, df.shape[1]
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    row_h = int(height / n_rows)
    for ri in range(n_rows):
        table.rows[ri].height = row_h
    for ci, col in enumerate(df.columns):
        set_cell(table.cell(0, ci), col, bold=True)
    for ri in range(df.shape[0]):
        for ci in range(n_cols):
            set_cell(table.cell(ri + 1, ci), df.iat[ri, ci])


def add_picture_fit(slide, img_path, left, top, max_w, max_h):
    with PILImage.open(img_path) as im:
        iw, ih = im.size
    ratio = iw / ih
    if ratio > max_w / max_h:
        w, h = max_w, max_w / ratio
    else:
        h, w = max_h, max_h * ratio
    slide.shapes.add_picture(str(img_path),
                              left + (max_w - w) / 2, top + (max_h - h) / 2, w, h)


# --------------------------------------------------------------------------
# Load data
# --------------------------------------------------------------------------
cont = pd.read_csv(CONTINUOUS)
cont[["mean", "sd", "median", "q25", "q75"]] = cont[["mean", "sd", "median", "q25", "q75"]].round(2)

cat = pd.read_csv(CATEGORICAL)
cat["pct"] = cat["pct"].round(1)

strat = pd.read_csv(STRATIFIED)
pair_cols  = [c for c in strat.columns if "_vs_" in c]
group_cols = [c for c in strat.columns if c not in ("label", "p_value", "q_value") + tuple(pair_cols)]
strat = strat[["label"] + group_cols + ["p_value", "q_value"] + pair_cols]
strat = strat.rename(columns={"label": "Variable", "p_value": "p (omnibus)", "q_value": "q (FDR)"})
strat = strat.rename(columns={c: c.replace("p_", "p ").replace("_vs_", " vs ") for c in pair_cols})

# --------------------------------------------------------------------------
# Build presentation
# --------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# Slide 1: clinical descriptive statistics
sl1 = add_slide(prs)
add_title(sl1, "Clinical descriptive statistics")
add_df_table(sl1, cont, Inches(0.3), Inches(1.0), Inches(6.3), Inches(6.2))
add_df_table(sl1, cat,  Inches(6.8), Inches(1.0), Inches(6.3), Inches(6.2))

# Slide 2: clinical descriptive statistics, stratified by target variable
sl2 = add_slide(prs)
add_title(sl2, "Clinical descriptive statistics — stratified")
add_df_table(sl2, strat, Inches(0.3), Inches(1.0), Inches(12.7), Inches(6.2))

# Slide 3: metabolomics summary heatmap
sl3 = add_slide(prs)
add_title(sl3, "Metabolomics — summary heatmap")
add_picture_fit(sl3, HEATMAP_METABOLOMICS, Inches(0.3), Inches(0.9), Inches(12.7), Inches(6.3))

# Slide 4: proteomics summary heatmap
sl4 = add_slide(prs)
add_title(sl4, "Proteomics — summary heatmap")
add_picture_fit(sl4, HEATMAP_PROTEOMICS, Inches(0.3), Inches(0.9), Inches(12.7), Inches(6.3))

# Slide 5: transcriptomics summary heatmap
sl5 = add_slide(prs)
add_title(sl5, "Transcriptomics — summary heatmap")
add_picture_fit(sl5, HEATMAP_TRANSCRIPTOMICS, Inches(0.3), Inches(0.9), Inches(12.7), Inches(6.3))

# --------------------------------------------------------------------------
# Save
# --------------------------------------------------------------------------
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Saved: {OUT}")
